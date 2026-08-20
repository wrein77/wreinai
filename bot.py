import os
import asyncio
import tempfile
import base64

from aiogram import Bot, Dispatcher, F
from aiogram.filters import CommandStart, Command
from aiogram.types import (
    Message,
    CallbackQuery,
    InlineKeyboardMarkup,
    InlineKeyboardButton
)

from groq import Groq

from pypdf import PdfReader
from docx import Document
import openpyxl
import pandas as pd


# ==================================================
# НАСТРОЙКИ
# ==================================================

BOT_TOKEN = os.getenv("8747215142:AAHQvxuno7sLoGJQr0-ryX82vUSKJDHOuZs")
GROQ_API_KEY = os.getenv("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")

TEXT_MODEL = "openai/gpt-oss-120b"
VISION_MODEL = "qwen/qwen3.6-27b"

if not BOT_TOKEN:
    raise RuntimeError("Не задан BOT_TOKEN")

if not GROQ_API_KEY:
    raise RuntimeError("Не задан GROQ_API_KEY")


bot = Bot(token=BOT_TOKEN)
dp = Dispatcher()
groq = Groq(api_key=GROQ_API_KEY)


# ==================================================
# ПАМЯТЬ
# ==================================================

memory = {}

MAX_MESSAGES = 12


SYSTEM_PROMPT = """
Ты — WREIN AI, персональный ИИ-ассистент в Telegram.

Отвечай естественно, понятно и без лишней воды.

Если пользователь пишет на русском — отвечай на русском.

Ты умеешь:
- отвечать на вопросы;
- объяснять сложные темы;
- помогать с программированием;
- помогать с учёбой;
- анализировать фотографии;
- распознавать текст на изображениях;
- работать с PDF;
- работать с DOCX;
- работать с TXT;
- работать с CSV;
- работать с XLSX.

Не утверждай, что сделал что-то, чего фактически не делал.

Если пользователь загрузил файл и задаёт вопрос по нему,
используй содержимое файла для ответа.
"""


# ==================================================
# КНОПКИ
# ==================================================

def main_keyboard():

    return InlineKeyboardMarkup(
        inline_keyboard=[
            [
                InlineKeyboardButton(
                    text="🧠 Новый чат",
                    callback_data="clear"
                )
            ],
            [
                InlineKeyboardButton(
                    text="ℹ️ Помощь",
                    callback_data="help"
                )
            ]
        ]
    )


# ==================================================
# START
# ==================================================

@dp.message(CommandStart())
async def start(message: Message):

    await message.answer(
        "🤖 <b>WREIN AI</b>\n\n"
        "Привет! Я твой ИИ-ассистент.\n\n"
        "📝 Напиши вопрос\n"
        "🖼 Отправь фотографию\n"
        "📄 Отправь PDF или DOCX\n"
        "📊 Отправь таблицу\n\n"
        "Я попробую разобраться.",
        reply_markup=main_keyboard(),
        parse_mode="HTML"
    )


# ==================================================
# HELP
# ==================================================

@dp.callback_query(F.data == "help")
async def help_callback(callback: CallbackQuery):

    await callback.message.answer(
        "ℹ️ <b>Что я умею</b>\n\n"
        "🧠 ИИ-чат\n"
        "🖼 Анализ фотографий\n"
        "🔤 OCR текста с изображений\n"
        "📄 PDF\n"
        "📘 DOCX\n"
        "📝 TXT\n"
        "📊 CSV\n"
        "📗 XLSX\n"
        "💻 Программирование\n"
        "📚 Учёба\n\n"
        "/clear — очистить память",
        parse_mode="HTML"
    )

    await callback.answer()


# ==================================================
# CLEAR
# ==================================================

@dp.message(Command("clear"))
async def clear_command(message: Message):

    memory.pop(message.from_user.id, None)

    await message.answer(
        "🧹 Память очищена.\n\n"
        "Начинаем новый разговор."
    )


@dp.callback_query(F.data == "clear")
async def clear_callback(callback: CallbackQuery):

    memory.pop(callback.from_user.id, None)

    await callback.message.answer(
        "🧹 <b>Новый чат создан.</b>\n\n"
        "Старый контекст очищен.",
        parse_mode="HTML"
    )

    await callback.answer("Готово")


# ==================================================
# GROQ TEXT
# ==================================================

def groq_text(messages):

    result = groq.chat.completions.create(
        model=TEXT_MODEL,
        messages=messages,
        temperature=0.7,
        max_completion_tokens=2048,
        stream=False
    )

    return result.choices[0].message.content


async def ask_ai(user_id, text):

    if user_id not in memory:

        memory[user_id] = [
            {
                "role": "system",
                "content": SYSTEM_PROMPT
            }
        ]

    memory[user_id].append(
        {
            "role": "user",
            "content": text
        }
    )

    memory[user_id] = (
        [memory[user_id][0]]
        + memory[user_id][-MAX_MESSAGES:]
    )

    answer = await asyncio.to_thread(
        groq_text,
        memory[user_id]
    )

    memory[user_id].append(
        {
            "role": "assistant",
            "content": answer
        }
    )

    return answer


# ==================================================
# ТЕКСТОВЫЕ СООБЩЕНИЯ
# ==================================================

@dp.message(F.text)
async def text_handler(message: Message):

    if message.text.startswith("/"):
        return

    await bot.send_chat_action(
        message.chat.id,
        "typing"
    )

    try:

        answer = await ask_ai(
            message.from_user.id,
            message.text
        )

        await send_long_message(
            message,
            answer
        )

    except Exception as e:

        print("TEXT ERROR:", e)

        await message.answer(
            "❌ Произошла ошибка при обращении к ИИ."
        )


# ==================================================
# ФОТО
# ==================================================

@dp.message(F.photo)
async def photo_handler(message: Message):

    await bot.send_chat_action(
        message.chat.id,
        "typing"
    )

    path = None

    try:

        photo = message.photo[-1]

        telegram_file = await bot.get_file(
            photo.file_id
        )

        with tempfile.NamedTemporaryFile(
            suffix=".jpg",
            delete=False
        ) as temp:

            path = temp.name

        await bot.download_file(
            telegram_file.file_path,
            path
        )

        with open(path, "rb") as image:

            encoded = base64.b64encode(
                image.read()
            ).decode("utf-8")

        question = message.caption

        if not question:

            question = (
                "Подробно проанализируй это изображение. "
                "Если на нём есть текст, распознай его."
            )

        result = await asyncio.to_thread(
            lambda: groq.chat.completions.create(
                model=VISION_MODEL,
                messages=[
                    {
                        "role": "system",
                        "content": SYSTEM_PROMPT
                    },
                    {
                        "role": "user",
                        "content": [
                            {
                                "type": "text",
                                "text": question
                            },
                            {
                                "type": "image_url",
                                "image_url": {
                                    "url":
                                        "data:image/jpeg;base64,"
                                        + encoded
                                }
                            }
                        ]
                    }
                ],
                temperature=0.4,
                max_completion_tokens=2048,
                stream=False
            )
        )

        answer = result.choices[0].message.content

        await send_long_message(
            message,
            answer
        )

    except Exception as e:

        print("IMAGE ERROR:", e)

        await message.answer(
            "❌ Не получилось обработать фотографию."
        )

    finally:

        if path and os.path.exists(path):
            os.remove(path)


# ==================================================
# СКАЧИВАНИЕ ФАЙЛА
# ==================================================

async def download_file(message):

    telegram_file = await bot.get_file(
        message.document.file_id
    )

    filename = message.document.file_name or "file"

    extension = os.path.splitext(
        filename
    )[1].lower()

    with tempfile.NamedTemporaryFile(
        suffix=extension,
        delete=False
    ) as temp:

        path = temp.name

    await bot.download_file(
        telegram_file.file_path,
        path
    )

    return path


# ==================================================
# PDF
# ==================================================

def read_pdf(path):

    reader = PdfReader(path)

    text = ""

    for page in reader.pages:

        page_text = page.extract_text()

        if page_text:
            text += page_text + "\n"

        if len(text) >= 50000:
            break

    return text[:50000]


# ==================================================
# DOCX
# ==================================================

def read_docx(path):

    document = Document(path)

    result = []

    for paragraph in document.paragraphs:

        if paragraph.text.strip():

            result.append(
                paragraph.text
            )

    return "\n".join(result)


# ==================================================
# TXT
# ==================================================

def read_txt(path):

    for encoding in (
        "utf-8",
        "cp1251",
        "latin-1"
    ):

        try:

            with open(
                path,
                "r",
                encoding=encoding
            ) as file:

                return file.read()[:50000]

        except UnicodeDecodeError:

            continue

    return ""


# ==================================================
# CSV
# ==================================================

def read_csv(path):

    dataframe = pd.read_csv(
        path,
        nrows=500
    )

    return dataframe.to_string()[:50000]


# ==================================================
# XLSX
# ==================================================

def read_xlsx(path):

    workbook = openpyxl.load_workbook(
        path,
        read_only=True,
        data_only=True
    )

    result = []

    for sheet in workbook.worksheets:

        result.append(
            f"\n--- {sheet.title} ---"
        )

        for row in sheet.iter_rows(
            values_only=True
        ):

            values = []

            for value in row:

                values.append(
                    str(value)
                    if value is not None
                    else ""
                )

            result.append(
                " | ".join(values)
            )

    return "\n".join(result)[:50000]


# ==================================================
# AI ДЛЯ ФАЙЛА
# ==================================================

def ask_file_ai(filename, text, question):

    prompt = f"""
Пользователь отправил файл:

Название: {filename}

Содержимое:

--- BEGIN FILE ---
{text}
--- END FILE ---

Запрос пользователя:

{question}

Ответь на русском языке.

Если вопрос касается содержимого файла,
используй именно предоставленный текст.
"""

    result = groq.chat.completions.create(
        model=TEXT_MODEL,
        messages=[
            {
                "role": "system",
                "content": SYSTEM_PROMPT
            },
            {
                "role": "user",
                "content": prompt
            }
        ],
        temperature=0.3,
        max_completion_tokens=2048,
        stream=False
    )

    return result.choices[0].message.content


# ==================================================
# ФАЙЛЫ
# ==================================================

@dp.message(F.document)
async def document_handler(message: Message):

    await bot.send_chat_action(
        message.chat.id,
        "typing"
    )

    path = None

    try:

        filename = (
            message.document.file_name
            or "file"
        )

        extension = os.path.splitext(
            filename
        )[1].lower()

        allowed = {
            ".pdf",
            ".txt",
            ".docx",
            ".csv",
            ".xlsx"
        }

        if extension not in allowed:

            await message.answer(
                "❌ Такой формат пока не поддерживается.\n\n"
                "Поддерживаются:\n"
                "📄 PDF\n"
                "📝 TXT\n"
                "📘 DOCX\n"
                "📊 CSV\n"
                "📗 XLSX"
            )

            return

        path = await download_file(
            message
        )

        if extension == ".pdf":

            text = read_pdf(path)

        elif extension == ".docx":

            text = read_docx(path)

        elif extension == ".txt":

            text = read_txt(path)

        elif extension == ".csv":

            text = read_csv(path)

        elif extension == ".xlsx":

            text = read_xlsx(path)

        else:

            text = ""

        if not text.strip():

            await message.answer(
                "❌ В файле не удалось найти текст."
            )

            return

        question = message.caption

        if not question:

            question = (
                "Кратко проанализируй файл "
                "и расскажи, что в нём находится."
            )

        answer = await asyncio.to_thread(
            ask_file_ai,
            filename,
            text,
            question
        )

        await send_long_message(
            message,
            answer
        )

    except Exception as e:

        print("FILE ERROR:", e)

        await message.answer(
            "❌ Не получилось обработать файл."
        )

    finally:

        if path and os.path.exists(path):

            os.remove(path)


# ==================================================
# ДЛИННЫЕ СООБЩЕНИЯ
# ==================================================

async def send_long_message(
    message,
    text
):

    if not text:

        text = "Не удалось получить ответ."

    for i in range(
        0,
        len(text),
        4000
    ):

        await message.answer(
            text[i:i + 4000]
        )


# ==================================================
# ЗАПУСК
# ==================================================

async def main():

    print("🤖 WREIN AI запущен!")

    await dp.start_polling(bot)


if __name__ == "__main__":

    asyncio.run(main())
